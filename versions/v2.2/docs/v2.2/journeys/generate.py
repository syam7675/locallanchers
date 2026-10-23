from pathlib import Path
from xml.sax.saxutils import escape
import textwrap
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import ParagraphStyle
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
ROOT=Path(__file__).resolve().parents[3];OUT=Path(__file__).resolve().parent
for name,file in [('Body','DejaVuSans.ttf'),('Bold','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
GREEN='173E36';BG='F5F3ED';MUTED='53645B';SAGE='DFE8D3';WHITE='FFFFFF'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
pdf=canvas.Canvas(str(OUT/'Local-Lanchers-v2.2-Journey-Presentation.pdf'),pagesize=(960,540))
slide=None;page=0

def box(x,y,w,h,color=WHITE):
 s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background()
 pdf.setFillColor(HexColor('#'+color));pdf.roundRect(x*72,540-(y+h)*72,w*72,h*72,10,stroke=0,fill=1)
def text(x,y,w,h,content,size=18,bold=False,color=GREEN):
 tf=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(6)
 style=ParagraphStyle('t',fontName='Bold' if bold else 'Body',fontSize=size,leading=size*1.22,textColor=HexColor('#'+color),spaceAfter=0)
 p=Paragraph(escape(content).replace('\n','<br/>'),style);_,height=p.wrap(w*72,h*72)
 if height>h*72+2:raise ValueError(('Overflow',content,height,h*72))
 p.drawOn(pdf,x*72,540-y*72-height)
def pic(path,x,y,w,h):
 im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
 slide.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
 pdf.drawImage(str(path),(x+(w-ww)/2)*72,540-(y+(h+hh)/2)*72,width=ww*72,height=hh*72,mask='auto')
def start(kicker,title,subtitle):
 global slide,page
 if page:pdf.showPage()
 page+=1;slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(BG);pdf.setFillColor(HexColor('#'+BG));pdf.rect(0,0,960,540,fill=1,stroke=0)
 text(.55,.35,10,.25,kicker,10,True,MUTED);text(.55,.9,12.1,1,title,32,True);text(.55,1.95,12.1,.65,subtitle,15,False,MUTED)
 text(.55,7.08,11,.2,'LOCAL LANCHERS  /  VERSION 2.2  /  CUSTOMER + ADMIN JOURNEYS',8,False,MUTED);text(12,7.05,.7,.25,f'{page:02d}',11,True)
def card(x,y,w,h,label,title,body):
 box(x,y,w,h);text(x+.22,y+.2,w-.44,.25,label,10,True,MUTED);text(x+.22,y+.65,w-.44,.8,title,20,True);text(x+.22,y+1.55,w-.44,h-1.7,body,14,False,MUTED)

start('THE PROJECT GOAL','Local businesses. A clearer path to growth.','A practical guide to what customers see, what admins do, and what comes next.')
pic(ROOT/'docs/v2.2/verification/home.png',6.15,2.65,6.6,4.1)
box(.55,2.8,5.3,3.8,GREEN);text(.85,3.1,4.7,.5,'DISCOVER → CONVERT → RETAIN',15,True,SAGE);text(.85,3.9,4.7,1.4,'The customer is a\nlocal business owner.',25,True,WHITE);text(.85,5.55,4.7,.7,'Current public site: an interactive preview.\nReal intake and accounts need persistent hosting.',13,False,WHITE)

start('01 / FIND YOUR WAY','One familiar navigation. Two working journeys.','The shared header serves visitors; workspace and operations serve different tasks.')
items=[('Home + Solutions','/  ·  /growth','Understand the offer and four growth pillars.'),('Industries + Plans','/industries  ·  /plans','Choose a sector and compare indicative options.'),('How it works + Audit','/how-it-works  ·  /audit','Understand delivery, then request a conversation.'),('Demo + Footer','/workspace  ·  /ops','Try workflows; admins use the private operations desk.')]
for i,(title,path,body) in enumerate(items):
 x=.55+(i%2)*6.2;y=2.8+(i//2)*1.9;box(x,y,5.95,1.65);text(x+.22,y+.18,5.5,.35,title,19,True);text(x+.22,y+.65,5.5,.25,path,11,True,MUTED);text(x+.22,y+1.05,5.5,.42,body,13,False,MUTED)

start('02 / CUSTOMER JOURNEY','From interest to an agreed growth plan.','The audit starts a human conversation. It does not automatically book or charge the customer.')
steps=[('01','Understand','Home → Solutions\nIdentify the main growth priority.'),('02','Choose a fit','Industry → Process → Plans\nCompare the approach and scope.'),('03','Request an audit','Business → Contact → Review\nSubmit when the backend is enabled.'),('04','Agree the work','Team follow-up → proposal\nConfirm scope, access and terms manually.')]
for i,(num,title,body) in enumerate(steps):card(.55+i*3.12,2.9,2.92,3.15,num,title,body)
box(.55,6.3,12.2,.5,SAGE);text(.75,6.43,11.8,.25,'A saved reference confirms receipt. Proposals, payments, scheduling and onboarding agreements remain outside the app.',11)

start('03 / WORKSPACE + ACTIVITY','Try the workflow. Understand the rhythm.','Growth overview · Enquiry desk · Review drafts · Visibility setup · Onboarding')
card(.55,2.85,5.85,3.7,'AVAILABLE ON THE PUBLIC PREVIEW','Demo profiles','No login needed. Choose “Act as”, perform practice actions, then inspect the heatmap. Filter by person, day or 30/90/365-day range.')
card(6.65,2.85,6.1,3.7,'FULL BACKEND WHEN ENABLED','My real activity','Create an account or sign in. Your successful logins and workspace interactions form a private history. Clear history or delete the account when needed.')
text(.65,6.65,12,.25,'Quiet 0  ·  Light 1–2  ·  Steady 3–5  ·  Active 6–9  ·  Busy 10+ actions per day. Counts describe app usage, not growth.',10,False,MUTED)

start('04 / ADMIN JOURNEY','Turn an audit request into a clear next step.','Open /ops and use the configured operations password. Customer accounts use a separate sign-in.')
steps=[('NEW','Review the request','Check business details, priorities and permission to reply.'),('REVIEWING','Prepare the audit','Review the opportunity and save useful internal notes.'),('CONTACTED','Follow up manually','Use an agreed channel outside the app. Record the outcome.'),('COMPLETED','Close the request','Save the final status and next step; export the queue if needed.')]
for i,(label,title,body) in enumerate(steps):card(.55+i*3.12,2.9,2.92,3.2,label,title,body)
box(.55,6.3,12.2,.5,SAGE);text(.75,6.43,11.8,.25,'The operations desk does not manage customer passwords, organization roles, subscriptions or all users’ heatmaps.',11)

start('05 / REACH THE GOAL','Launch the service. Then grow the software.','A sensible sequence from the current preview to a working growth operation.')
card(.55,2.9,3.9,3.7,'FIRST','Enable real intake','Persistent storage, HTTPS and secure team access. Test saved requests and account sign-in. Finalize privacy and support details.')
card(4.7,2.9,3.9,3.7,'NEXT','Run a small pilot','Manually deliver audit → proposal → onboarding → review. Prove the workflow with a focused group of local businesses.')
card(8.85,2.9,3.9,3.7,'THEN','Connect and measure','Connect authorized providers. Measure qualified enquiries, attended bookings, repeat visits and revenue from real source data.')
pdf.save();prs.save(OUT/'Local-Lanchers-v2.2-Journey-Presentation.pptx')

# The brief is intentionally two pages; the presentation is a separate six-slide artifact.
styles={
 'title':ParagraphStyle('title',fontName='Bold',fontSize=23,leading=28,textColor=HexColor('#'+GREEN),spaceAfter=12),
 'h':ParagraphStyle('h',fontName='Bold',fontSize=12,leading=16,textColor=HexColor('#'+GREEN),spaceBefore=12,spaceAfter=6),
 'body':ParagraphStyle('body',fontName='Body',fontSize=9,leading=13,textColor=HexColor('#'+GREEN),spaceAfter=7),
 'small':ParagraphStyle('small',fontName='Body',fontSize=8,leading=11,textColor=HexColor('#'+MUTED),spaceAfter=5)}
story=[]
def para(s,style='body'):story.append(Paragraph(s,styles[style]))
def header(c,d):
 c.setFillColor(HexColor('#'+GREEN));c.rect(0,809,595,33,fill=1,stroke=0);c.setFillColor(HexColor('#'+WHITE));c.setFont('Bold',9);c.drawString(40,821,'LOCAL LANCHERS   /   VERSION 2.2');c.setFillColor(HexColor('#'+MUTED));c.setFont('Body',8);c.drawString(40,24,'Customer + admin journey brief  •  23 October 2026');c.drawRightString(555,24,str(d.page))
para('How the website works','title')
para('For local business owners and the operations team.','small')
para('<b>Goal:</b> help a local business get discovered, turn enquiries into customers, retain them and review growth. The website’s customer is the business owner—not the end consumer booking a clinic or salon visit.')
para('<b>Today:</b> the Vercel site is a public design/demo preview. Real audit submissions, individual sign-in and private operations require the full backend with persistent storage. External messaging, booking, payment, Google and AI services are not connected.')
para('Frontend and navigation','h')
para('A shared header/footer surrounds the pages. Desktop dropdowns and the mobile menu reach the same destinations. Forms validate before calling the server. Practice records stay in the browser; enabled real account activity is stored separately on the server.')
rows=[['Go to','Purpose'],['Home / Solutions','Understand the offer and four growth pillars.'],['Industries / How it works','Choose dental, salon or fitness; understand delivery.'],['Plans','Compare indicative options and calculator scenarios.'],['Get a growth audit','Business details → contact/permission → review.'],['Explore the demo','Growth overview, enquiries, reviews, visibility, onboarding.'],['Footer','Privacy, build readiness and Team operations desk (/ops).']]
t=Table([[Paragraph(escape(x),styles['small']) for x in row] for row in rows],colWidths=[156,359]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#'+SAGE)),('ROWBACKGROUNDS',(0,1),(-1,-1),[HexColor('#FFFFFF'),HexColor('#F5F3ED')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.append(t)
para('Customer: the path to an agreed plan','h')
for s in ['<b>1. Understand.</b> Home → Solutions. Identify discovery, conversion, retention or measurement as the priority.','<b>2. Find your fit.</b> Industry → How it works → Plans. Options are indicative; there is no checkout.','<b>3. Request an audit when enabled.</b> Complete the three steps and submit. A saved reference means received, not an appointment booked.','<b>4. Agree the work.</b> The team follows up manually, reviews opportunities and agrees scope, access, delivery and commercial terms outside the app.','<b>5. Explore the workspace.</b> Try enquiry states, review drafts and checklists. These remain practice workflows, including after account sign-in.']:para(s)
story.append(PageBreak())
para('Login, activity and operations','title')
para('Customer login and the activity calendar','h')
para('<b>Demo:</b> open Explore the demo → Growth overview → Demo profiles. No login is needed. Choose “Act as”, perform actions in the existing sections, then return to the heatmap. Filter by profile and 30/90/365 days; select a day to view events.')
para('<b>Real account, on an enabled backend:</b> My real activity → Create an account. Enter name, email, a password of at least 12 characters and agreement to activity recording. Existing users choose Sign in. Successful logins and workspace interactions appear in your own private history. Sign out, Clear history and Delete account are available. Email verification and password recovery are not implemented.')
para('<b>Read the heatmap:</b> Quiet 0 · Light 1–2 · Steady 3–5 · Active 6–9 · Busy 10+ actions/day. Dates use India time. These are feature-use counts—not evidence of bookings, revenue, productivity or retention.')
para('Admin: the path from request to follow-through','h')
for s in ['<b>1. Prepare hosting.</b> The operator enables persistent storage and HTTPS, sets PUBLIC_ORIGIN and a strong OPS_PASSWORD, and optionally ENABLE_USER_ACCOUNTS=true. Test persistence before inviting customers.','<b>2. Sign in at /ops.</b> Use the configured operations password. This is separate from customer email/password login; it is not a full role-based administration console.','<b>3. Review.</b> Search/filter requests, check business details and contact permission, set New → Reviewing, and save notes.','<b>4. Follow up.</b> Contact the owner outside the app. Set Contacted after outreach and record agreed next steps. Use Completed when the audit request is handled—not as proof of growth.','<b>5. Maintain the queue.</b> Save updates and export a filtered CSV if needed. The desk does not manage customer passwords, organization roles, subscriptions or all users’ heatmaps.']:para(s)
para('What to do next to achieve the business goal','h')
para('<b>First:</b> enable durable intake and accounts. <b>Next:</b> validate the manual audit → proposal → onboarding → review process with a small pilot. <b>Then:</b> connect the providers needed for delivery and measure qualified enquiries, attended bookings, repeat visits and revenue using authorized source data.')
para('Live preview: <link href="https://local-lanchers-v22.vercel.app" color="#173E36">local-lanchers-v22.vercel.app</link>','small')
para('Prepared from the V2.2 source and release notes. Higgsfield was unavailable in this session; the companion presentation was generated locally.','small')
SimpleDocTemplate(str(OUT/'Local-Lanchers-v2.2-Customer-Admin-Guide.pdf'),pagesize=(595,842),leftMargin=40,rightMargin=40,topMargin=52,bottomMargin=40,title='Local Lanchers 2.2 — Customer and Admin Guide',author='Local Lanchers').build(story,onFirstPage=header,onLaterPages=header)
print('Created guide PDF, presentation PDF and editable PPTX in',OUT)
